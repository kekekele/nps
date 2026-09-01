from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import stats

from .core import (
    SCORE_CODES,
    benjamini_hochberg,
    clean_text,
    cramers_v,
    information_value,
    is_subquestion_hit,
    month_start,
    mutual_information,
    normalize_identifier,
    parse_score,
    parse_subquestion_header,
    segment_scores,
    standardized_mean_difference,
)

SURVEY_SHEET = "测评明细"
TECH_HEADER_SHEETS = {
    "基础指标",
    "近半年指标",
    "资费变更",
    "限速与加包",
    "触点轨迹",
    "AI感知总结",
}
SCORE_PREFIXES = {code: f"{code}." for code in SCORE_CODES}
MONTHLY_COLUMNS = [
    "month",
    "id_no",
    "phone",
    "customer_type",
    "broadband",
    "has_broadband",
    "other_phone_broadband",
    "sim_slot",
    "ported_in",
    "queried_port_code",
    "open_date",
    "arpu",
    "voice_total",
    "data_total",
    "voice_overage_fee",
    "data_overage_fee",
    "status_code",
    "broadband_fee",
    "included_data",
    "used_data",
    "included_voice",
    "used_voice",
    "roaming_outside_province",
]
BASIC_COLUMNS = [
    "phone",
    "customer_type",
    "phone_209",
    "is_household",
    "city",
    "district",
    "service_center",
    "gender",
    "age",
    "open_date",
    "tenure_months",
    "global_level",
    "device_vendor",
    "device_model",
    "device_type",
    "main_plan_fee",
    "scene",
    "enterprise",
    "phone_280",
    "is_13n",
]
HOUSEHOLD_COLUMNS = [
    "phone",
    "customer_type",
    "broadband",
    "bandwidth",
    "broadband_area",
    "open_date",
    "broadband_fee",
    "is_primary_payer",
    "broadband_usage",
]


@dataclass
class PipelineContext:
    workbook: Path
    output_root: Path
    batch_id: str
    version: str
    generated_at: str = field(default_factory=lambda: datetime.now(UTC).isoformat())
    files: list[dict[str, Any]] = field(default_factory=list)
    rejects: list[dict[str, Any]] = field(default_factory=list)

    @property
    def clean_dir(self) -> Path:
        return self.output_root / "clean" / self.batch_id

    @property
    def analysis_dir(self) -> Path:
        return self.output_root / "analysis" / self.version

    @property
    def dictionary_dir(self) -> Path:
        return self.output_root / "dictionary"

    @property
    def quality_dir(self) -> Path:
        return self.output_root / "quality" / self.batch_id

    @property
    def manifest_dir(self) -> Path:
        return self.output_root / "manifest"

    def prepare(self) -> None:
        for directory in (
            self.clean_dir,
            self.analysis_dir,
            self.dictionary_dir,
            self.quality_dir,
            self.manifest_dir,
        ):
            directory.mkdir(parents=True, exist_ok=True)

    def write_csv(self, frame: pd.DataFrame, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        frame.to_csv(
            path, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d %H:%M:%S"
        )
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        self.files.append(
            {
                "path": path.relative_to(self.output_root).as_posix(),
                "rows": int(len(frame)),
                "columns": int(len(frame.columns)),
                "sha256": digest,
            }
        )

    def reject(
        self,
        sheet: str,
        row: int | None,
        field_name: str,
        raw_value: Any,
        rule_code: str,
    ) -> None:
        self.rejects.append(
            {
                "source_sheet": sheet,
                "source_row_no": row,
                "field_name": field_name,
                "raw_value": clean_text(raw_value),
                "rule_code": rule_code,
                "source_batch_id": self.batch_id,
            }
        )


def _read_sheet(workbook: Path, sheet_name: str) -> pd.DataFrame:
    excel = load_workbook(workbook, read_only=True, data_only=True)
    try:
        worksheet = excel[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        if sheet_name == SURVEY_SHEET:
            next(rows)
        headers = list(next(rows))
        if sheet_name in TECH_HEADER_SHEETS:
            next(rows)
        return pd.DataFrame.from_records(rows, columns=headers)
    finally:
        excel.close()


def _rename_by_position(frame: pd.DataFrame, names: list[str]) -> pd.DataFrame:
    if len(frame.columns) != len(names):
        raise ValueError(f"Expected {len(names)} columns, found {len(frame.columns)}")
    result = frame.copy()
    result.columns = names
    return result


def _clean_scalar(value: Any) -> Any:
    if isinstance(value, str):
        return clean_text(value)
    return None if pd.isna(value) else value


def _clean_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    object_columns = result.select_dtypes(include=["object"]).columns
    for column in object_columns:
        result[column] = result[column].map(_clean_scalar)
    return result


def _yes_no(value: Any) -> int | None:
    text = clean_text(value)
    if text in {"是", "Y", "YES", "1", "TRUE"}:
        return 1
    if text in {"否", "N", "NO", "0", "FALSE"}:
        return 0
    return None


def _date_value(value: Any) -> pd.Timestamp:
    text = normalize_identifier(value)
    if text is None:
        return pd.NaT
    if re.fullmatch(r"\d{8}", text):
        return pd.to_datetime(text, format="%Y%m%d", errors="coerce")
    return pd.to_datetime(value, errors="coerce")


def _score_column(headers: list[str], code: str) -> str:
    prefix = SCORE_PREFIXES[code]
    matches = [header for header in headers if header.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one score column for {code}, found {matches}")
    return matches[0]


def build_survey_outputs(ctx: PipelineContext) -> dict[str, pd.DataFrame]:
    raw = _read_sheet(ctx.workbook, SURVEY_SHEET)
    raw.columns = [
        clean_text(column) or f"unnamed_{index + 1}"
        for index, column in enumerate(raw.columns)
    ]
    raw["source_row_no"] = np.arange(3, len(raw) + 3)
    headers = list(raw.columns)
    score_columns = {code: _score_column(headers, code) for code in SCORE_CODES}
    survey_id_column = "调研流水号"
    phone_column = "手机号码"

    duplicate_ids = (
        raw[survey_id_column].map(normalize_identifier).duplicated(keep=False)
    )
    for _, row in raw.loc[duplicate_ids].iterrows():
        ctx.reject(
            SURVEY_SHEET,
            int(row["source_row_no"]),
            survey_id_column,
            row[survey_id_column],
            "duplicate_survey_id",
        )

    records: list[dict[str, Any]] = []
    score_records: list[dict[str, Any]] = []
    for _, row in raw.iterrows():
        survey_id = normalize_identifier(row[survey_id_column])
        customer_key = normalize_identifier(row[phone_column])
        if survey_id is None:
            ctx.reject(
                SURVEY_SHEET,
                int(row["source_row_no"]),
                survey_id_column,
                row[survey_id_column],
                "missing_survey_id",
            )
            continue
        values = [row[column] for column in score_columns.values()]
        segment, valid_count, low_count = segment_scores(values)
        records.append(
            {
                "survey_id": survey_id,
                "customer_key": customer_key,
                "survey_month": month_start(row["调研时间"]),
                "survey_scene": clean_text(row["调研业务场景名称"]),
                "city": clean_text(row["地市名称"]),
                "district": clean_text(row["区县名称"]),
                "service_center": clean_text(row["服务中心名称"]),
                "is_fttr": _yes_no(row["是否FTTR"]),
                "is_bandwidth_ge_500m": _yes_no(row["是否带宽大于等于500M"]),
                "is_traffic_customer": _yes_no(row["是否流量运营客户"]),
                "valid_score_count": valid_count,
                "low_score_count": low_count,
                "segment_label": segment,
                "source_row_no": int(row["source_row_no"]),
                "source_batch_id": ctx.batch_id,
            }
        )
        for code, column in score_columns.items():
            score, invalid_reason = parse_score(row[column])
            score_records.append(
                {
                    "survey_id": survey_id,
                    "customer_key": customer_key,
                    "question_code": code,
                    "raw_value": clean_text(row[column]),
                    "score_value": score,
                    "is_valid": int(score is not None),
                    "is_low_score": int(score is not None and score <= 6),
                    "invalid_reason": invalid_reason,
                }
            )
            if invalid_reason not in {None, "missing_or_sentinel"}:
                ctx.reject(
                    SURVEY_SHEET,
                    int(row["source_row_no"]),
                    code,
                    row[column],
                    invalid_reason,
                )

    survey = pd.DataFrame(records).sort_values(["survey_id", "source_row_no"])
    survey = survey.drop_duplicates("survey_id", keep="last").reset_index(drop=True)
    valid_ids = set(survey["survey_id"])
    scores = pd.DataFrame(score_records)
    scores = scores[scores["survey_id"].isin(valid_ids)].drop_duplicates(
        ["survey_id", "question_code"]
    )

    question_dictionary = pd.DataFrame(
        [
            {
                "question_code": code,
                "question_text": score_columns[code],
                "include_in_segment": 1,
                "min_score": 0,
                "max_score": 10,
            }
            for code in SCORE_CODES
        ]
    )

    option_columns: list[tuple[str, dict[str, str]]] = []
    for header in headers:
        parsed = parse_subquestion_header(header)
        if parsed:
            option_columns.append((header, parsed))
    option_dictionary = pd.DataFrame(
        [parsed for _, parsed in option_columns]
    ).drop_duplicates()

    hit_records: list[dict[str, Any]] = []
    for _, row in raw.iterrows():
        survey_id = normalize_identifier(row[survey_id_column])
        if survey_id not in valid_ids:
            continue
        for column, option in option_columns:
            hit, error = is_subquestion_hit(row[column], option["option_code"])
            if hit:
                hit_records.append(
                    {
                        "survey_id": survey_id,
                        "question_code": option["parent_question_code"],
                        "subquestion_code": option["subquestion_code"],
                        "option_code": option["option_code"],
                        "option_text": option["option_text"],
                        "hit_flag": 1,
                        "raw_value": clean_text(row[column]),
                    }
                )
            elif error:
                ctx.reject(
                    SURVEY_SHEET, int(row["source_row_no"]), column, row[column], error
                )
    hits = pd.DataFrame(
        hit_records,
        columns=[
            "survey_id",
            "question_code",
            "subquestion_code",
            "option_code",
            "option_text",
            "hit_flag",
            "raw_value",
        ],
    )

    q99_column = next(column for column in headers if column.startswith("Q99."))
    voices: list[dict[str, Any]] = []
    for _, row in raw.iterrows():
        survey_id = normalize_identifier(row[survey_id_column])
        text = clean_text(row[q99_column])
        valid = text not in {None, "-1"} and bool(
            re.search(r"[\w\u4e00-\u9fff]", text or "")
        )
        if survey_id in valid_ids and valid:
            clean_voice = re.sub(r"\s+", " ", text or "").strip()
            voices.append(
                {
                    "voice_id": hashlib.sha256(f"{survey_id}|Q99".encode()).hexdigest(),
                    "survey_id": survey_id,
                    "customer_key": normalize_identifier(row[phone_column]),
                    "voice_type": "Q99",
                    "voice_text_raw": text,
                    "voice_text_clean": clean_voice,
                    "language": "zh",
                    "text_length": len(clean_voice),
                    "is_valid": 1,
                    "sensitive_info_status": "source_masking_retained",
                }
            )
    voice = pd.DataFrame(
        voices,
        columns=[
            "voice_id",
            "survey_id",
            "customer_key",
            "voice_type",
            "voice_text_raw",
            "voice_text_clean",
            "language",
            "text_length",
            "is_valid",
            "sensitive_info_status",
        ],
    )

    ctx.write_csv(survey, ctx.clean_dir / "fact_survey.csv")
    ctx.write_csv(scores, ctx.clean_dir / "fact_survey_score.csv")
    ctx.write_csv(question_dictionary, ctx.dictionary_dir / "dim_question.csv")
    ctx.write_csv(option_dictionary, ctx.dictionary_dir / "dim_subquestion_option.csv")
    ctx.write_csv(hits, ctx.clean_dir / "bridge_survey_subquestion_hit.csv")
    ctx.write_csv(voice, ctx.clean_dir / "fact_customer_voice.csv")
    return {"survey": survey, "scores": scores, "hits": hits, "voice": voice}


def build_customer_outputs(ctx: PipelineContext) -> dict[str, pd.DataFrame]:
    basic = _clean_frame(
        _rename_by_position(_read_sheet(ctx.workbook, "基础指标"), BASIC_COLUMNS)
    )
    basic["customer_key"] = basic["phone"].map(normalize_identifier)
    basic["phone_209_key"] = basic["phone_209"].map(normalize_identifier)
    basic["phone_280_key"] = basic["phone_280"].map(normalize_identifier)
    basic["is_household"] = basic["is_household"].map(_yes_no)
    basic["is_13n"] = basic["is_13n"].map(_yes_no)
    basic["age"] = pd.to_numeric(basic["age"], errors="coerce")
    basic["tenure_months"] = pd.to_numeric(basic["tenure_months"], errors="coerce")
    basic["main_plan_fee"] = pd.to_numeric(basic["main_plan_fee"], errors="coerce")
    basic["open_date"] = basic["open_date"].map(_date_value)
    basic["source_batch_id"] = ctx.batch_id
    basic = basic.drop(columns=["phone", "phone_209", "phone_280"])
    basic = basic.dropna(subset=["customer_key"])
    basic = basic.sort_values(
        ["customer_key", "customer_type"], na_position="last"
    ).drop_duplicates("customer_key")

    household = _clean_frame(
        _rename_by_position(_read_sheet(ctx.workbook, "家庭指标"), HOUSEHOLD_COLUMNS)
    )
    household["customer_key"] = household["phone"].map(normalize_identifier)
    household["household_key"] = household["broadband"].map(normalize_identifier)
    household["broadband_key"] = household["household_key"]
    household["is_primary_payer"] = household["is_primary_payer"].map(_yes_no)
    household["member_role"] = np.where(
        household["is_primary_payer"].eq(1), "payer", "member"
    )
    household["open_date"] = household["open_date"].map(_date_value)
    for column in ("bandwidth", "broadband_fee", "broadband_usage"):
        household[column] = pd.to_numeric(household[column], errors="coerce")
    household["source_batch_id"] = ctx.batch_id

    dim_household = (
        household.dropna(subset=["household_key"])[
            [
                "household_key",
                "broadband_key",
                "bandwidth",
                "broadband_area",
                "open_date",
                "broadband_fee",
                "broadband_usage",
                "source_batch_id",
            ]
        ]
        .sort_values("open_date")
        .drop_duplicates("household_key", keep="last")
    )
    members = household.dropna(subset=["household_key", "customer_key"])[
        [
            "household_key",
            "customer_key",
            "member_role",
            "is_primary_payer",
            "source_batch_id",
        ]
    ].drop_duplicates(["household_key", "customer_key"])
    members["effective_start"] = pd.NaT
    members["effective_end"] = pd.NaT
    members["relation_source"] = "shared_broadband"
    members["confidence_score"] = 1.0

    ctx.write_csv(basic, ctx.clean_dir / "dim_customer.csv")
    ctx.write_csv(dim_household, ctx.clean_dir / "dim_household.csv")
    ctx.write_csv(members, ctx.clean_dir / "bridge_household_member.csv")
    return {"customer": basic, "household": dim_household, "members": members}


def build_monthly_outputs(ctx: PipelineContext) -> dict[str, pd.DataFrame]:
    monthly = _clean_frame(
        _rename_by_position(_read_sheet(ctx.workbook, "近半年指标"), MONTHLY_COLUMNS)
    )
    monthly["customer_key"] = monthly["phone"].map(normalize_identifier)
    monthly["broadband_key"] = monthly["broadband"].map(normalize_identifier)
    monthly["month"] = monthly["month"].map(month_start)
    monthly["open_date"] = monthly["open_date"].map(_date_value)
    monthly["has_broadband"] = monthly["has_broadband"].map(_yes_no)
    monthly["ported_in"] = monthly["ported_in"].map(_yes_no)
    monthly["queried_port_code"] = monthly["queried_port_code"].map(_yes_no)
    monthly["roaming_outside_province"] = monthly["roaming_outside_province"].map(
        _yes_no
    )
    numeric_metrics = [
        "arpu",
        "voice_total",
        "data_total",
        "voice_overage_fee",
        "data_overage_fee",
        "broadband_fee",
        "included_data",
        "used_data",
        "included_voice",
        "used_voice",
    ]
    for column in numeric_metrics:
        monthly[column] = pd.to_numeric(monthly[column], errors="coerce")
    monthly = monthly.dropna(subset=["customer_key", "month"])
    key = ["customer_key", "month", "customer_type", "broadband_key"]
    monthly = monthly.sort_values(key, na_position="last").drop_duplicates(
        key, keep="last"
    )

    monthly_long = monthly.melt(
        id_vars=["customer_key", "month"],
        value_vars=numeric_metrics,
        var_name="metric_code",
        value_name="metric_value",
    )
    monthly_long["missing_reason"] = np.where(
        monthly_long["metric_value"].isna(), "source_missing", None
    )
    ctx.write_csv(monthly_long, ctx.clean_dir / "fact_customer_monthly_metric.csv")
    trajectory_columns = ["customer_key", "month", *numeric_metrics]
    monthly_for_analysis = monthly[trajectory_columns].copy()
    del monthly_long, monthly
    return {"monthly": monthly_for_analysis}


EVENT_SPECS = {
    "业务订购": {
        "time": 5,
        "event_type": "subscription",
        "subtype": 3,
        "detail": 4,
        "start_row": 2,
    },
    "投诉明细": {
        "time": 2,
        "event_type": "complaint",
        "subtype": 3,
        "detail": 5,
        "start_row": 2,
    },
    "资费变更": {
        "time": 10,
        "event_type": "plan_change",
        "subtype": 9,
        "detail": 6,
        "start_row": 3,
    },
    "限速与加包": {
        "time": 6,
        "event_type": "speed_package",
        "subtype": 0,
        "detail": 8,
        "start_row": 3,
    },
    "狼号": {
        "time": 2,
        "event_type": "special_contact",
        "subtype": 3,
        "detail": 1,
        "start_row": 2,
    },
    "触点轨迹": {
        "time": 4,
        "event_type": "contact",
        "subtype": 5,
        "detail": 13,
        "start_row": 3,
    },
}


def build_event_output(ctx: PipelineContext) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for sheet_name, spec in EVENT_SPECS.items():
        raw = _read_sheet(ctx.workbook, sheet_name)
        phone_index = (
            1 if sheet_name == "资费变更" else (2 if sheet_name == "限速与加包" else 0)
        )
        rows: list[dict[str, Any]] = []
        for offset, values in enumerate(
            raw.itertuples(index=False, name=None), start=spec["start_row"]
        ):
            customer_key = normalize_identifier(values[phone_index])
            event_time = _date_value(values[spec["time"]])
            if customer_key is None:
                ctx.reject(
                    sheet_name,
                    offset,
                    "phone",
                    values[phone_index],
                    "missing_customer_identifier",
                )
                continue
            if pd.isna(event_time):
                ctx.reject(
                    sheet_name,
                    offset,
                    "event_time",
                    values[spec["time"]],
                    "invalid_event_time",
                )
                continue
            detail = clean_text(values[spec["detail"]])
            event_identity = f"{sheet_name}|{customer_key}|{event_time}|{clean_text(values[spec['subtype']])}|{detail}"
            rows.append(
                {
                    "event_id": hashlib.sha256(
                        event_identity.encode("utf-8")
                    ).hexdigest(),
                    "customer_key": customer_key,
                    "event_time": event_time,
                    "event_type": spec["event_type"],
                    "event_subtype": clean_text(values[spec["subtype"]]),
                    "detail_ref": (
                        hashlib.sha256((detail or "").encode("utf-8")).hexdigest()
                        if detail
                        else None
                    ),
                    "source_sheet": sheet_name,
                    "source_row_no": offset,
                }
            )
        frames.append(pd.DataFrame(rows))
    events = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    events = events.sort_values("event_time").drop_duplicates("event_id")
    ctx.write_csv(events, ctx.clean_dir / "fact_customer_event.csv")
    return events


def build_analysis_sample(
    ctx: PipelineContext,
    survey: pd.DataFrame,
    customer: pd.DataFrame,
    members: pd.DataFrame,
    monthly: pd.DataFrame,
    events: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    household_features = members.groupby("customer_key", as_index=False).agg(
        household_count=("household_key", "nunique"),
        is_primary_payer=("is_primary_payer", "max"),
    )
    household_size = (
        members.groupby("household_key")["customer_key"]
        .nunique()
        .rename("household_size")
    )
    member_size = (
        members.merge(household_size, on="household_key")
        .groupby("customer_key", as_index=False)["household_size"]
        .max()
    )

    sample = survey.merge(
        customer, on="customer_key", how="left", suffixes=("_survey", "_customer")
    )
    for column in ("city", "district", "service_center"):
        survey_column = f"{column}_survey"
        customer_column = f"{column}_customer"
        if survey_column in sample.columns:
            sample[column] = sample[survey_column].combine_first(
                sample.get(customer_column)
            )
            sample = sample.drop(
                columns=[
                    name
                    for name in (survey_column, customer_column)
                    if name in sample.columns
                ]
            )
    sample = sample.merge(household_features, on="customer_key", how="left").merge(
        member_size, on="customer_key", how="left"
    )
    sample["as_of_date"] = sample["survey_month"]
    sample["feature_window"] = "M-6_to_M-1"

    trajectory = survey[
        ["survey_id", "customer_key", "survey_month", "segment_label"]
    ].merge(monthly, on="customer_key", how="left")
    trajectory["month_offset"] = (
        (trajectory["month"].dt.year - trajectory["survey_month"].dt.year) * 12
        + trajectory["month"].dt.month
        - trajectory["survey_month"].dt.month
    )
    trajectory = trajectory[
        trajectory["month_offset"].between(-6, -1, inclusive="both")
    ]
    numeric_metrics = [
        "arpu",
        "voice_total",
        "data_total",
        "voice_overage_fee",
        "data_overage_fee",
        "broadband_fee",
        "included_data",
        "used_data",
        "included_voice",
        "used_voice",
    ]
    trajectory = trajectory.groupby(
        [
            "survey_id",
            "customer_key",
            "survey_month",
            "segment_label",
            "month",
            "month_offset",
        ],
        as_index=False,
    )[numeric_metrics].mean()
    feature_frames: list[pd.DataFrame] = []
    for metric in numeric_metrics:
        grouped = trajectory.groupby("survey_id").agg(
            **{
                f"{metric}_mean_6m": (metric, "mean"),
                f"{metric}_std_6m": (metric, "std"),
                f"{metric}_min_6m": (metric, "min"),
                f"{metric}_max_6m": (metric, "max"),
                f"{metric}_valid_months": (metric, "count"),
            }
        )
        pivot = trajectory.pivot_table(
            index="survey_id", columns="month_offset", values=metric, aggfunc="last"
        )
        missing = pd.Series(np.nan, index=pivot.index, dtype=float)
        last_value = pivot.get(-1, missing)
        first_value = pivot.get(-6, missing)
        grouped[f"{metric}_last"] = last_value
        grouped[f"{metric}_change_6m"] = last_value - first_value
        feature_frames.append(grouped)
    if feature_frames:
        trajectory_features = pd.concat(feature_frames, axis=1).reset_index()
        sample = sample.merge(trajectory_features, on="survey_id", how="left")

    survey_events = survey[["survey_id", "customer_key", "survey_month"]].merge(
        events, on="customer_key", how="left"
    )
    survey_events["days_before_survey"] = (
        survey_events["survey_month"] - survey_events["event_time"].dt.normalize()
    ).dt.days
    survey_events = survey_events[
        survey_events["days_before_survey"].between(1, 183, inclusive="both")
    ]
    if not survey_events.empty:
        event_features = survey_events.groupby("survey_id").agg(
            event_count_6m=("event_id", "count"),
            event_type_count_6m=("event_type", "nunique"),
            days_since_last_event=("days_before_survey", "min"),
        )
        for window in (30, 90):
            counts = (
                survey_events[survey_events["days_before_survey"].le(window)]
                .groupby("survey_id")["event_id"]
                .count()
            )
            event_features[f"event_count_{window}d"] = counts
        type_counts = survey_events.pivot_table(
            index="survey_id",
            columns="event_type",
            values="event_id",
            aggfunc="count",
            fill_value=0,
        )
        type_counts.columns = [
            f"event_{column}_count_6m" for column in type_counts.columns
        ]
        event_features = event_features.join(type_counts).reset_index()
        sample = sample.merge(event_features, on="survey_id", how="left")
    ctx.write_csv(sample, ctx.analysis_dir / "survey_analysis_sample.csv")
    return sample, trajectory


def _cohort_distributions(sample: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    dimensions = ["segment_label", "survey_month", "survey_scene", "city", "district"]
    for dimension in dimensions:
        values = sample[dimension].fillna("<MISSING>")
        counts = values.value_counts(dropna=False)
        for value, count in counts.items():
            rows.append(
                {
                    "dimension": dimension,
                    "dimension_value": value,
                    "survey_count": int(count),
                    "survey_share": count / len(sample) if len(sample) else np.nan,
                    "unique_customer_count": int(
                        sample.loc[values.eq(value), "customer_key"].nunique()
                    ),
                }
            )
    return pd.DataFrame(rows)


def _feature_result(
    feature: pd.Series, labels: pd.Series, name: str, source: str
) -> dict[str, Any] | None:
    valid_label = labels.isin(["low", "non_low"])
    x = feature.loc[valid_label]
    y = labels.loc[valid_label]
    if x.notna().sum() < 20 or y.nunique() < 2 or x.nunique(dropna=True) < 2:
        return None
    numeric = pd.to_numeric(x, errors="coerce")
    non_missing = max(int(x.notna().sum()), 1)
    is_numeric = numeric.notna().sum() / non_missing >= 0.8 and numeric.nunique() >= 5
    analysis_x = x
    if is_numeric:
        try:
            analysis_x = pd.qcut(numeric, q=10, duplicates="drop").astype(str)
        except ValueError:
            analysis_x = numeric.astype(str)
        analysis_x = analysis_x.where(numeric.notna(), "<MISSING>")
    else:
        category = x.fillna("<MISSING>").astype(str)
        minimum_count = max(20, int(len(category) * 0.01))
        common = category.value_counts()[lambda counts: counts >= minimum_count].index
        analysis_x = category.where(category.isin(common), "<OTHER>")
    low = x[y.eq("low")]
    high = x[y.eq("non_low")]
    if is_numeric:
        low_num = pd.to_numeric(low, errors="coerce").dropna()
        high_num = pd.to_numeric(high, errors="coerce").dropna()
        p_value = (
            stats.mannwhitneyu(low_num, high_num, alternative="two-sided").pvalue
            if len(low_num) and len(high_num)
            else np.nan
        )
        effect = standardized_mean_difference(low_num, high_num)
        effect_name = "standardized_mean_difference"
        low_stat = low_num.median() if len(low_num) else np.nan
        high_stat = high_num.median() if len(high_num) else np.nan
        direction = (
            "low_higher"
            if low_stat > high_stat
            else "low_lower" if low_stat < high_stat else "equal"
        )
    else:
        grouped = analysis_x.fillna("<MISSING>").astype(str)
        contingency = pd.crosstab(grouped, y)
        p_value = (
            stats.chi2_contingency(contingency)[1]
            if min(contingency.shape) >= 2
            else np.nan
        )
        effect = cramers_v(contingency)
        effect_name = "cramers_v"
        rates = pd.crosstab(grouped, y, normalize="index").get(
            "low", pd.Series(dtype=float)
        )
        low_stat = rates.max() if not rates.empty else np.nan
        high_stat = rates.min() if not rates.empty else np.nan
        direction = str(rates.idxmax()) if not rates.empty else None
    overall_low_rate = y.eq("low").mean()
    grouped_rates = pd.crosstab(
        analysis_x.fillna("<MISSING>").astype(str), y, normalize="index"
    ).get("low", pd.Series(dtype=float))
    max_lift = (
        grouped_rates.max() / overall_low_rate
        if overall_low_rate and not grouped_rates.empty
        else np.nan
    )
    return {
        "feature_name": name,
        "source": source,
        "feature_type": "numeric" if is_numeric else "categorical",
        "sample_count": int(len(x)),
        "valid_count": int(x.notna().sum()),
        "coverage_rate": float(x.notna().mean()),
        "missing_rate": float(x.isna().mean()),
        "low_statistic": low_stat,
        "non_low_statistic": high_stat,
        "p_value": p_value,
        "effect_name": effect_name,
        "effect_size": effect,
        "information_value": information_value(analysis_x, y),
        "mutual_information": mutual_information(analysis_x, y),
        "max_low_rate_lift": max_lift,
        "direction_or_top_risk_group": direction,
    }


def analyze_features(
    sample: pd.DataFrame, trajectory_columns: set[str]
) -> tuple[pd.DataFrame, pd.DataFrame]:
    excluded = {
        "survey_id",
        "customer_key",
        "segment_label",
        "source_row_no",
        "source_batch_id",
        "survey_month",
        "as_of_date",
        "feature_window",
        "valid_score_count",
        "low_score_count",
    }
    rows: list[dict[str, Any]] = []
    for column in sample.columns:
        if (
            column in excluded
            or column.endswith("_key")
            or column.startswith("phone_")
            or column.startswith("source_")
            or pd.api.types.is_datetime64_any_dtype(sample[column])
        ):
            continue
        result = _feature_result(
            sample[column],
            sample["segment_label"],
            column,
            "trajectory" if column in trajectory_columns else "static",
        )
        if result:
            rows.append(result)
    results = pd.DataFrame(rows)
    if results.empty:
        return results, results
    results["fdr"] = benjamini_hochberg(results["p_value"])
    effect_component = results["effect_size"].abs().clip(upper=1).fillna(0)
    iv_component = results["information_value"].clip(upper=0.5).fillna(0) / 0.5
    coverage_component = results["coverage_rate"].fillna(0)
    significance_component = (1 - results["fdr"].fillna(1)).clip(lower=0)
    results["importance_score"] = (
        0.35 * effect_component
        + 0.25 * iv_component
        + 0.20 * coverage_component
        + 0.20 * significance_component
    )
    results["importance_level"] = pd.cut(
        results["importance_score"],
        bins=[-np.inf, 0.35, 0.55, np.inf],
        labels=["low", "medium", "high"],
    ).astype(str)
    results = results.sort_values(
        ["importance_score", "feature_name"], ascending=[False, True]
    )
    return (
        results[results["source"].eq("static")],
        results[results["source"].eq("trajectory")],
    )


def build_feature_groups(
    sample: pd.DataFrame, feature_names: list[str]
) -> pd.DataFrame:
    numeric = sample[
        [column for column in feature_names if column in sample.columns]
    ].apply(pd.to_numeric, errors="coerce")
    numeric = numeric.loc[:, numeric.notna().sum().ge(20)]
    if numeric.shape[1] < 2:
        return pd.DataFrame(
            columns=["feature_name", "related_feature", "association", "group_rule"]
        )
    correlation = numeric.corr(method="spearman").abs()
    rows: list[dict[str, Any]] = []
    for index, feature in enumerate(correlation.columns):
        for related in correlation.columns[index + 1 :]:
            association = correlation.loc[feature, related]
            if pd.notna(association) and association >= 0.8:
                rows.append(
                    {
                        "feature_name": feature,
                        "related_feature": related,
                        "association": association,
                        "group_rule": "absolute_spearman_ge_0.8",
                    }
                )
    return pd.DataFrame(
        rows, columns=["feature_name", "related_feature", "association", "group_rule"]
    )


def build_analysis_outputs(
    ctx: PipelineContext,
    survey_data: dict[str, pd.DataFrame],
    customer_data: dict[str, pd.DataFrame],
    monthly_data: dict[str, pd.DataFrame],
    events: pd.DataFrame,
) -> None:
    survey = survey_data["survey"]
    sample, trajectory = build_analysis_sample(
        ctx,
        survey,
        customer_data["customer"],
        customer_data["members"],
        monthly_data["monthly"],
        events,
    )
    ctx.write_csv(
        _cohort_distributions(sample), ctx.analysis_dir / "cohort_distribution.csv"
    )

    score_distribution = (
        survey_data["scores"]
        .groupby(
            ["question_code", "score_value", "is_low_score"],
            dropna=False,
            as_index=False,
        )
        .agg(
            response_count=("survey_id", "count"),
            unique_customer_count=("customer_key", "nunique"),
        )
    )
    ctx.write_csv(score_distribution, ctx.analysis_dir / "score_distribution.csv")

    hits = survey_data["hits"].merge(
        survey[["survey_id", "segment_label"]], on="survey_id", how="left"
    )
    root_cause = (
        hits.groupby(
            [
                "segment_label",
                "question_code",
                "subquestion_code",
                "option_code",
                "option_text",
            ],
            as_index=False,
        ).agg(hit_count=("survey_id", "count"), survey_count=("survey_id", "nunique"))
        if not hits.empty
        else hits
    )
    ctx.write_csv(root_cause, ctx.analysis_dir / "root_cause_summary.csv")

    trajectory_distribution = (
        trajectory.melt(
            id_vars=["survey_id", "segment_label", "month_offset"],
            value_vars=[
                "arpu",
                "voice_total",
                "data_total",
                "voice_overage_fee",
                "data_overage_fee",
                "broadband_fee",
                "included_data",
                "used_data",
                "included_voice",
                "used_voice",
            ],
            var_name="metric_code",
            value_name="metric_value",
        )
        .groupby(["segment_label", "month_offset", "metric_code"], as_index=False)
        .agg(
            sample_count=("metric_value", "count"),
            mean=("metric_value", "mean"),
            median=("metric_value", "median"),
            std=("metric_value", "std"),
            q25=("metric_value", lambda values: values.quantile(0.25)),
            q75=("metric_value", lambda values: values.quantile(0.75)),
        )
    )
    ctx.write_csv(
        trajectory_distribution, ctx.analysis_dir / "trajectory_distribution.csv"
    )

    trajectory_columns = {
        column
        for column in sample.columns
        if re.search(r"_(mean|std|min|max|last|change|valid)_?", column)
        or column.startswith("event_")
        or column == "days_since_last_event"
    }
    static_features, trajectory_features = analyze_features(sample, trajectory_columns)
    ctx.write_csv(static_features, ctx.analysis_dir / "important_feature_static.csv")
    ctx.write_csv(
        trajectory_features, ctx.analysis_dir / "important_feature_trajectory.csv"
    )
    feature_groups = build_feature_groups(
        sample,
        pd.concat(
            [static_features["feature_name"], trajectory_features["feature_name"]]
        ).tolist(),
    )
    ctx.write_csv(feature_groups, ctx.analysis_dir / "important_feature_group.csv")

    event_summary = (
        events.assign(
            event_month=events["event_time"].dt.to_period("M").dt.to_timestamp()
        )
        .groupby(
            ["event_month", "event_type", "event_subtype"], dropna=False, as_index=False
        )
        .agg(
            event_count=("event_id", "count"),
            unique_customer_count=("customer_key", "nunique"),
        )
    )
    ctx.write_csv(event_summary, ctx.analysis_dir / "event_distribution.csv")


def write_quality_outputs(ctx: PipelineContext, survey: pd.DataFrame) -> None:
    rejects = pd.DataFrame(
        ctx.rejects,
        columns=[
            "source_sheet",
            "source_row_no",
            "field_name",
            "raw_value",
            "rule_code",
            "source_batch_id",
        ],
    )
    ctx.write_csv(rejects, ctx.quality_dir / "dq_reject_record.csv")
    quality = pd.DataFrame(
        [
            {"metric": "survey_rows", "value": len(survey)},
            {"metric": "unique_survey_ids", "value": survey["survey_id"].nunique()},
            {"metric": "unique_customers", "value": survey["customer_key"].nunique()},
            {"metric": "low_surveys", "value": survey["segment_label"].eq("low").sum()},
            {
                "metric": "non_low_surveys",
                "value": survey["segment_label"].eq("non_low").sum(),
            },
            {
                "metric": "unlabeled_surveys",
                "value": survey["segment_label"].eq("unlabeled").sum(),
            },
            {"metric": "rejected_values", "value": len(rejects)},
        ]
    )
    ctx.write_csv(quality, ctx.quality_dir / "data_quality_summary.csv")


def write_manifest(ctx: PipelineContext) -> Path:
    workbook_hash = hashlib.sha256(ctx.workbook.read_bytes()).hexdigest()
    manifest = {
        "batch_id": ctx.batch_id,
        "version": ctx.version,
        "generated_at": ctx.generated_at,
        "source_file": ctx.workbook.name,
        "source_sha256": workbook_hash,
        "files": ctx.files,
    }
    path = ctx.manifest_dir / f"{ctx.batch_id}.json"
    path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return path


def run_pipeline(
    workbook: Path, output_root: Path, batch_id: str, version: str
) -> Path:
    ctx = PipelineContext(workbook.resolve(), output_root.resolve(), batch_id, version)
    ctx.prepare()
    print("[1/7] Cleaning survey data", flush=True)
    survey_data = build_survey_outputs(ctx)
    print("[2/7] Cleaning customer and household data", flush=True)
    customer_data = build_customer_outputs(ctx)
    print("[3/7] Cleaning monthly metrics", flush=True)
    monthly_data = build_monthly_outputs(ctx)
    print("[4/7] Cleaning event data", flush=True)
    events = build_event_output(ctx)
    print("[5/7] Building statistical analysis outputs", flush=True)
    build_analysis_outputs(ctx, survey_data, customer_data, monthly_data, events)
    print("[6/7] Writing data quality outputs", flush=True)
    write_quality_outputs(ctx, survey_data["survey"])
    print("[7/7] Writing manifest", flush=True)
    return write_manifest(ctx)
