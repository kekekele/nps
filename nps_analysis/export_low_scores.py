from __future__ import annotations

import argparse
import csv
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCORE_QUESTIONS = {"Q1", "Q2", "Q3", "Q4", "Q5", "Q98", "N1"}
REASON_QUESTIONS = {"Q1", "Q2", "Q3", "Q4", "Q5", "Q98"}
SCORE_RE = re.compile(r"^(Q(?:[1-5]|98)|N1)\s*\.")
REASON_RE = re.compile(r"^(Q(?:[1-5]|98))-1\s*\.")
OPTION_TEXT_RE = re.compile(r"（可多选）\s*\d+\s*(.+?)\s*$")
EMPTY_VALUES = {"", "-1", r"\N", "None", "nan"}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_low_score(value: Any) -> bool:
    try:
        return 1 <= float(text(value)) <= 6
    except ValueError:
        return False


def reason_text(header: str, value: Any) -> str:
    if "填写内容" in header:
        return text(value)
    match = OPTION_TEXT_RE.search(header)
    return match.group(1).strip() if match else header.strip()


def export(input_path: Path, output_path: Path) -> None:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet_name = next(
        (name for name in ("测评明细", "测评明细数据") if name in workbook.sheetnames),
        None,
    )
    if sheet_name is None:
        raise ValueError(f"找不到测评明细 Sheet；实际 Sheet：{workbook.sheetnames}")

    rows = workbook[sheet_name].iter_rows(values_only=True)
    next(rows, None)  # 第1行是分组说明
    headers = [text(value) for value in next(rows)]  # 第2行是字段名
    phone_index = headers.index("手机号码")

    score_columns: dict[str, int] = {}
    reason_columns: dict[str, list[int]] = {q: [] for q in REASON_QUESTIONS}
    for index, header in enumerate(headers):
        if match := SCORE_RE.match(header):
            score_columns[match.group(1)] = index
        elif match := REASON_RE.match(header):
            reason_columns[match.group(1)].append(index)

    missing = SCORE_QUESTIONS - score_columns.keys()
    if missing:
        raise ValueError(f"缺少评分列：{', '.join(sorted(missing))}")

    users: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in rows:
        phone_id = text(row[phone_index])
        if not phone_id:
            continue
        user = users.setdefault(
            phone_id, {"is_low_score": 0, "reasons": OrderedDict()}
        )
        for question, score_index in score_columns.items():
            if not is_low_score(row[score_index]):
                continue
            user["is_low_score"] = 1
            # N1 没有 N1-1 原因题，只参与低分标记。
            for reason_index in reason_columns.get(question, []):
                value = text(row[reason_index])
                if value in EMPTY_VALUES:
                    continue
                reason = reason_text(headers[reason_index], value)
                if reason and reason not in EMPTY_VALUES:
                    user["reasons"][reason] = None

    with output_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["phone_id", "is_low_score", "reason"])
        writer.writeheader()
        for phone_id, user in users.items():
            writer.writerow({
                "phone_id": phone_id,
                "is_low_score": user["is_low_score"],
                "reason": ";".join(user["reasons"]),
            })

    workbook.close()
    print(f"完成：{len(users)} 个用户，输出到 {output_path}")


def main() -> None:
    base = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="按手机号汇总测评低分及原因")
    parser.add_argument("--input", type=Path, default=base / "指标汇总.xlsx")
    parser.add_argument("--output", type=Path, default=base / "用户低分汇总.csv")
    args = parser.parse_args()
    export(args.input.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
