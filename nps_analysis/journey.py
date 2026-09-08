from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .core import clean_text, normalize_identifier, parse_business_datetime
from .journey_config import (
    BusinessRule,
    DEFAULT_JOURNEY_CONFIG,
    EventSourceConfig,
    JourneyConfig,
    TextValueRule,
)


def _indexes(sheet: Worksheet, row: int) -> dict[str, int]:
    return {
        clean_text(cell.value) or "": index
        for index, cell in enumerate(sheet[row])
        if clean_text(cell.value)
    }


def _value(row: tuple[Any, ...], indexes: dict[str, int], column: str) -> Any:
    index = indexes.get(clean_text(column) or column)
    if index is None:
        raise ValueError(f"Column {column!r} was not found")
    return row[index]


def _time(value: Any) -> tuple[datetime | None, bool]:
    parsed = parse_business_datetime(value)
    if getattr(parsed, "to_pydatetime", None) is None:
        return None, False
    moment = parsed.to_pydatetime()
    has_time = (
        isinstance(value, datetime)
        or (isinstance(value, str) and bool(re.search(r"\d{1,2}:\d{2}", value)))
        or (
            str(value).replace(".0", "").isdigit()
            and len(str(value).replace(".0", "")) > 8
        )
    )
    return moment, has_time


def _format_time(moment: datetime, has_time: bool) -> str:
    return moment.strftime("%Y-%m-%d %H:%M:%S" if has_time else "%Y-%m-%d")


def _in_month_range(moment: datetime, month_range: tuple[str, ...]) -> bool:
    if len(month_range) != 2 or not month_range[0] or not month_range[1]:
        return True
    current = moment.strftime("%Y-%m")
    return month_range[0] <= current <= month_range[1]


def _event_id(parts: list[str]) -> int:
    digest = hashlib.sha256("|".join(parts).encode("utf-8")).digest()
    return int.from_bytes(digest[:8], "big") & ((1 << 63) - 1)


def _event(
    phone: str,
    moment: datetime,
    has_time: bool,
    action: str,
    business: str,
    intent: str,
    source: str,
    row_no: int,
    key: str,
) -> dict[str, Any]:
    return {
        "phone_id": phone,
        "event_id": _event_id(
            [source, str(row_no), phone, _format_time(moment, has_time), key]
        ),
        "event_time": _format_time(moment, has_time),
        "_moment": moment,
        "_has_time": has_time,
        "_source": source,
        "action": action,
        "business": business,
        "intent": intent,
        "source": "",
        "raw_text": "",
        "complaint_handling_satisfaction": None,
    }


def _matches_pattern(value: str, pattern: str, match_type: str) -> bool:
    if match_type == "contains":
        return pattern in value
    return bool(re.search(pattern, value))


def _classify_business(text: str, rules: tuple[BusinessRule, ...], default: str) -> str:
    for rule in rules:
        if re.search(rule.pattern, text, flags=re.IGNORECASE):
            return rule.business
    return default


def _configured_month_range(
    config: JourneyConfig, source: EventSourceConfig
) -> tuple[str, ...]:
    if not source.month_range_name:
        return ()
    return getattr(config, source.month_range_name, ())


def _field_values(
    row: tuple[Any, ...], indexes: dict[str, int], source: EventSourceConfig
) -> dict[str, Any]:
    return {field: _value(row, indexes, field) for field in source.input_fields}


def _passes_filters(field_values: dict[str, Any], source: EventSourceConfig) -> bool:
    for row_filter in source.row_filters:
        text = clean_text(field_values.get(row_filter.column)) or ""
        matched = _matches_pattern(text, row_filter.pattern, row_filter.match_type)
        if row_filter.negate:
            matched = not matched
        if not matched:
            return False
    return True


def _binding_value(value: Any, transform: str) -> Any:
    if transform != "numeric":
        return clean_text(value) or ""
    text = clean_text(value)
    if not text or text == r"\N":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _assign_output_field(event: dict[str, Any], target: str, value: Any) -> None:
    if value in (None, ""):
        return
    if target == "business":
        event["business"] = value
    elif target == "intent":
        event["intent"] = value
    else:
        event[target] = value


def _resolve_business(
    field_values: dict[str, Any], source: EventSourceConfig, config: JourneyConfig
) -> str:
    event = {
        "business": source.field_template.business,
        "intent": source.field_template.intent,
    }
    for binding in source.output_bindings:
        raw_value = field_values.get(binding.column)
        if binding.transform == "business_category":
            value = _classify_business(
                clean_text(raw_value) or "",
                config.business_rules,
                config.default_business,
            )
        else:
            value = _binding_value(raw_value, binding.transform)
        _assign_output_field(event, binding.target, value)
    return event["business"]


def _resolve_intent(field_values: dict[str, Any], source: EventSourceConfig) -> str:
    if source.intent_column:
        text = clean_text(field_values.get(source.intent_column)) or ""
        for rule in source.intent_rules:
            if _matches_pattern(text, rule.pattern, rule.match_type):
                return rule.value
    return source.default_intent


def _resolve_reference(
    token: str,
    event: dict[str, Any],
    field_values: dict[str, Any],
) -> Any:
    if token == "$phone_id":
        return event["phone_id"]
    if token == "$moment":
        return event["_moment"]
    if token == "$action":
        return event["action"]
    if token == "$business":
        return event["business"]
    if token == "$intent":
        return event["intent"]
    return field_values.get(token)


def _build_business_key(
    source: EventSourceConfig,
    event: dict[str, Any],
    field_values: dict[str, Any],
    row_no: int,
) -> str:
    if not source.business_key_fields:
        return str(row_no)
    return "|".join(
        clean_text(_resolve_reference(token, event, field_values)) or ""
        for token in source.business_key_fields
    )


def _build_dedupe_key(
    source: EventSourceConfig,
    event: dict[str, Any],
    field_values: dict[str, Any],
) -> tuple[Any, ...] | None:
    if not source.dedupe_fields:
        return None
    return tuple(
        _resolve_reference(token, event, field_values) for token in source.dedupe_fields
    )


def _apply_text_value_rules(
    current_value: str,
    default_value: str,
    action: str,
    text: str,
    rules: tuple[TextValueRule, ...],
) -> str:
    if current_value != default_value:
        return current_value
    for rule in rules:
        if rule.apply_to and action not in rule.apply_to:
            continue
        if re.search(rule.pattern, text) and not (
            rule.negation_pattern and re.search(rule.negation_pattern, text)
        ):
            return rule.value
    return current_value


def _joined_text(texts: list[Any]) -> str:
    return " ".join(clean_text(value) or "" for value in texts)


def _apply_business_text_rules(
    event: dict[str, Any], texts: list[Any], config: JourneyConfig
) -> None:
    text = _joined_text(texts)
    event["business"] = _apply_text_value_rules(
        event["business"],
        config.default_business,
        event["action"],
        text,
        config.business_text_rules,
    )


def _apply_intent_text_rules(
    event: dict[str, Any], texts: list[Any], config: JourneyConfig
) -> None:
    text = _joined_text(texts)
    event["intent"] = _apply_text_value_rules(
        event["intent"],
        config.default_intent,
        event["action"],
        text,
        config.intent_text_rules,
    )


def _extract_source_events(
    sheet: Worksheet, source: EventSourceConfig, config: JourneyConfig
) -> list[dict[str, Any]]:
    indexes = _indexes(sheet, source.sheet.sheet.header_row)
    events = []
    for row_no, row in enumerate(
        sheet.iter_rows(min_row=source.sheet.sheet.data_start_row, values_only=True),
        source.sheet.sheet.data_start_row,
    ):
        phone = normalize_identifier(_value(row, indexes, source.sheet.phone_column))
        moment, has_time = _time(_value(row, indexes, source.sheet.time_column))
        if not phone or not moment:
            continue
        if not _in_month_range(moment, _configured_month_range(config, source)):
            continue
        field_values = _field_values(row, indexes, source)
        if not _passes_filters(field_values, source):
            continue
        event = _event(
            phone,
            moment,
            has_time,
            source.action,
            _resolve_business(field_values, source, config),
            _resolve_intent(field_values, source),
            source.source_name,
            row_no,
            str(row_no),
        )
        for binding in source.output_bindings:
            if binding.target not in {"business", "intent"}:
                _assign_output_field(
                    event,
                    binding.target,
                    _binding_value(
                        field_values.get(binding.column), binding.transform
                    ),
                )
        event["event_id"] = _event_id(
            [
                source.source_name,
                str(row_no),
                phone,
                _format_time(moment, has_time),
                _build_business_key(source, event, field_values, row_no),
            ]
        )
        if source.business_text_rule_fields:
            _apply_business_text_rules(
                event,
                [
                    field_values.get(column)
                    for column in source.business_text_rule_fields
                ],
                config,
            )
        if source.intent_text_rule_fields:
            _apply_intent_text_rules(
                event,
                [field_values.get(column) for column in source.intent_text_rule_fields],
                config,
            )
        dedupe_key = _build_dedupe_key(source, event, field_values)
        if dedupe_key is not None:
            event["_dedupe_key"] = dedupe_key
        events.append(event)
    return events


def _merge_event(target: dict[str, Any], incoming: dict[str, Any]) -> None:
    if target["business"] == "其他" and incoming["business"] != "其他":
        target["business"] = incoming["business"]
    if target["intent"] == "其他" and incoming["intent"] != "其他":
        target["intent"] = incoming["intent"]


def build_journeys(
    workbook: str | Path, config: JourneyConfig = DEFAULT_JOURNEY_CONFIG
) -> dict[str, list[dict[str, Any]]]:
    excel = load_workbook(Path(workbook), read_only=True, data_only=True)
    try:
        events = [
            event
            for source in config.event_sources
            if source.sheet.sheet.name in excel.sheetnames
            for event in _extract_source_events(
                excel[source.sheet.sheet.name], source, config
            )
        ]
        unique_events = {}
        for event in events:
            key = event.get("_dedupe_key", (event["_source"], event["event_id"]))
            existing = unique_events.get(key)
            if existing is None:
                unique_events[key] = event
            else:
                _merge_event(existing, event)
        by_phone: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for event in unique_events.values():
            by_phone[event["phone_id"]].append(event)
        for phone, journey in by_phone.items():
            journey.sort(
                key=lambda event: (
                    event["_moment"],
                    not event["_has_time"],
                    event["event_id"],
                )
            )
            for event in journey:
                event.pop("phone_id")
                event.pop("_moment")
                event.pop("_has_time")
                event.pop("_source")
                event.pop("_dedupe_key", None)
        return by_phone
    finally:
        excel.close()
