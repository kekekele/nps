from __future__ import annotations

import argparse
import json
from collections import defaultdict
from collections.abc import Iterable, Iterator
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .core import clean_text, normalize_identifier
from .profile_config import (
    DEFAULT_PROFILE_CONFIG,
    FieldSpec,
    SheetSpec,
    UserProfileConfig,
)
from .journey import build_journeys

TRUE_VALUES = {"是", "Y", "YES", "1", "TRUE"}
FALSE_VALUES = {"否", "N", "NO", "0", "FALSE"}


def _number_value(value: Any) -> int | float:
    if value is None:
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() else value
    text = clean_text(value)
    if text is None:
        return 0
    try:
        number = float(text)
    except ValueError:
        return 0
    return int(number) if number.is_integer() else number


def _bool_value(value: Any) -> int:
    text = (clean_text(value) or "").upper()
    if text in TRUE_VALUES:
        return 1
    if text in FALSE_VALUES or text == "":
        return 0
    return 0


def _clean_value(value: Any, value_type: str = "text") -> Any:
    if value_type == "number":
        return _number_value(value)
    if value_type == "bool":
        return _bool_value(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return clean_text(value) or ""
    return "" if value is None else value


def _column_indexes(worksheet: Worksheet, spec: SheetSpec) -> dict[str, list[int]]:
    indexes: dict[str, list[int]] = defaultdict(list)
    for index, cell in enumerate(worksheet[spec.header_row]):
        name = clean_text(cell.value)
        if name:
            indexes[name].append(index)
    return indexes


def _index(indexes: dict[str, list[int]], name: str, occurrence: int = 1) -> int:
    normalized_name = clean_text(name) or name
    matches = indexes.get(normalized_name, [])
    if len(matches) < occurrence:
        raise ValueError(f"Column {name!r} occurrence {occurrence} was not found")
    return matches[occurrence - 1]


def _project(
    row: tuple[Any, ...], indexes: dict[str, list[int]], fields: Iterable[FieldSpec]
) -> dict[str, Any]:
    return {
        field.key: _clean_value(
            row[_index(indexes, field.source, field.occurrence)], field.value_type
        )
        for field in fields
    }


def _rows(worksheet: Worksheet, spec: SheetSpec) -> Iterator[tuple[Any, ...]]:
    yield from worksheet.iter_rows(min_row=spec.data_start_row, values_only=True)


def _vendor_category(value: Any, config: UserProfileConfig) -> str:
    vendor = (clean_text(value) or "").upper()
    if not vendor:
        return ""
    for category, aliases in config.common_vendor_aliases.items():
        if any(alias.upper() in vendor for alias in aliases):
            return category
    return "其他"


def build_user_profiles(
    workbook: str | Path, config: UserProfileConfig = DEFAULT_PROFILE_CONFIG
) -> Iterator[dict[str, Any]]:
    excel = load_workbook(Path(workbook), read_only=True, data_only=True)
    try:
        journeys = build_journeys(workbook)
        survey = excel[config.survey_sheet.name]
        survey_indexes = _column_indexes(survey, config.survey_sheet)
        survey_phone_index = _index(survey_indexes, config.survey_phone_column)
        survey_by_phone = {
            normalize_identifier(row[survey_phone_index]): _project(
                row, survey_indexes, config.survey_fields
            )
            for row in _rows(survey, config.survey_sheet)
            if normalize_identifier(row[survey_phone_index])
        }

        household = excel[config.household_sheet.name]
        household_indexes = _column_indexes(household, config.household_sheet)
        broadband_index = _index(household_indexes, config.household_broadband_column)
        payer_index = _index(household_indexes, config.household_primary_payer_column)
        household_by_broadband = {
            normalize_identifier(row[broadband_index]): _project(
                row, household_indexes, config.household_fields
            )
            for row in _rows(household, config.household_sheet)
            if normalize_identifier(row[broadband_index])
            and _bool_value(row[payer_index]) == 1
        }

        monthly = excel[config.monthly_sheet.name]
        monthly_indexes = _column_indexes(monthly, config.monthly_sheet)
        monthly_phone_index = _index(monthly_indexes, config.monthly_phone_column)
        month_index = _index(monthly_indexes, config.monthly_month_column)
        monthly_by_phone: dict[str, list[tuple[Any, dict[str, Any]]]] = defaultdict(
            list
        )
        for row in _rows(monthly, config.monthly_sheet):
            phone = normalize_identifier(row[monthly_phone_index])
            if phone:
                monthly_by_phone[phone].append(
                    (
                        _clean_value(row[month_index]),
                        _project(row, monthly_indexes, config.monthly_fields),
                    )
                )

        basic = excel[config.basic_sheet.name]
        basic_indexes = _column_indexes(basic, config.basic_sheet)
        basic_phone_index = _index(basic_indexes, config.basic_phone_column)
        household_key_index = _index(basic_indexes, config.basic_household_column)
        seen: set[str] = set()
        for row in _rows(basic, config.basic_sheet):
            phone = normalize_identifier(row[basic_phone_index])
            if not phone or phone in seen:
                continue
            seen.add(phone)
            basic_info = _project(row, basic_indexes, config.basic_fields)
            basic_info["term_factory_category"] = _vendor_category(
                basic_info.get("term_factory"), config
            )
            basic_info.update(survey_by_phone.get(phone, {}))
            household_key = normalize_identifier(row[household_key_index])
            basic_info.update(household_by_broadband.get(household_key, {}))

            monthly_rows = sorted(
                monthly_by_phone.get(phone, []), key=lambda item: str(item[0])
            )
            recent_metrics = {"month": [month for month, _ in monthly_rows]}
            for field in config.monthly_fields:
                recent_metrics[field.key] = [
                    values[field.key] for _, values in monthly_rows
                ]

            yield {
                "phone_id": phone,
                "basic_info": basic_info,
                "recent_metrics": recent_metrics,
                "user_journey": journeys.get(phone, []),
            }
    finally:
        excel.close()


def export_user_profiles(
    workbook: str | Path,
    output: str | Path,
    config: UserProfileConfig = DEFAULT_PROFILE_CONFIG,
) -> int:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with output_path.open("w", encoding="utf-8") as stream:
        for profile in build_user_profiles(workbook, config):
            stream.write(json.dumps(profile, ensure_ascii=False) + "\n")
            count += 1
    return count


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export one JSON object per mobile user."
    )
    parser.add_argument("--input", type=Path, default=Path("doc/指标汇总.xlsx"))
    parser.add_argument("--output", type=Path, default=Path("data/user_profiles.jsonl"))
    args = parser.parse_args()
    count = export_user_profiles(args.input, args.output)
    print(f"Exported {count} user profiles to {args.output}")


if __name__ == "__main__":
    main()
