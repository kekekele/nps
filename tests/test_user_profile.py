import json

from openpyxl import Workbook

from nps_analysis.profile_config import DEFAULT_PROFILE_CONFIG
from nps_analysis.user_profile import build_user_profiles, export_user_profiles


def _add_sheet(workbook, spec, headers, rows, description_row=False):
    sheet = workbook.create_sheet(spec.name)
    for column, value in enumerate(headers, 1):
        sheet.cell(spec.header_row, column, value)
    for row_number, row in enumerate(rows, spec.data_start_row):
        for column, value in enumerate(row, 1):
            sheet.cell(row_number, column, value)


def _workbook(tmp_path):
    config = DEFAULT_PROFILE_CONFIG
    workbook = Workbook()
    workbook.remove(workbook.active)
    basic_headers = [config.basic_phone_column] + [
        field.source for field in config.basic_fields
    ]
    basic_row = ["13800000000"] + [r"\N"] * len(config.basic_fields)
    basic_row[basic_headers.index("209号码")] = "028-123"
    basic_row[basic_headers.index("终端厂商")] = "Apple Inc."
    _add_sheet(
        workbook, config.basic_sheet, basic_headers, [basic_row], description_row=True
    )

    survey_headers = [config.survey_phone_column] + [
        field.source for field in config.survey_fields
    ]
    _add_sheet(
        workbook,
        config.survey_sheet,
        survey_headers,
        [["13800000000", "是", "否", "是"]],
    )

    household_headers = [
        config.household_broadband_column,
        config.household_primary_payer_column,
        *[field.source for field in config.household_fields],
    ]
    _add_sheet(
        workbook,
        config.household_sheet,
        household_headers,
        [
            ["028-123", "否", 100, "A", 20200101, 50, 10],
            ["028-123", "是", 500, "B", 20210101, 80, 20],
        ],
    )

    monthly_headers = [config.monthly_phone_column, config.monthly_month_column]
    monthly_headers.extend(field.source for field in config.monthly_fields)
    monthly_rows = [
        ["13800000000", 202603, 30, 31, 32, 33, 34, 35, 36, 37, 38, 39, 40, "是"],
        ["13800000000", 202602, r"\N", 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, "否"],
    ]
    _add_sheet(
        workbook,
        config.monthly_sheet,
        monthly_headers,
        monthly_rows,
        description_row=True,
    )
    path = tmp_path / "profiles.xlsx"
    workbook.save(path)
    return path


def test_build_user_profiles_joins_and_sorts_source_data(tmp_path):
    profile = next(build_user_profiles(_workbook(tmp_path)))

    assert profile["phone_id"] == "13800000000"
    assert profile["basic_info"]["customer_type"] == ""
    assert profile["basic_info"]["age"] == 0
    assert profile["basic_info"]["term_factory_category"] == "苹果"
    assert profile["basic_info"]["is_fttr"] == 1
    assert profile["basic_info"]["broadband_bandwidth"] == 500
    assert profile["basic_info"]["broadband_usage"] == 20
    assert profile["recent_metrics"]["month"] == [202602, 202603]
    assert profile["recent_metrics"]["arpu"] == [0, 30]
    assert profile["recent_metrics"]["mou"] == [22, 32]
    assert profile["recent_metrics"]["gprs_used_v"] == [28, 38]
    assert profile["recent_metrics"]["roaming_outside_province"] == [0, 1]
    assert profile["user_journey"] == []


def test_export_user_profiles_writes_json_lines(tmp_path):
    output = tmp_path / "profiles.jsonl"
    count = export_user_profiles(_workbook(tmp_path), output)

    assert count == 1
    assert json.loads(output.read_text(encoding="utf-8"))["phone_id"] == "13800000000"
