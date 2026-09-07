from dataclasses import replace

from openpyxl import Workbook

from nps_analysis.journey import build_journeys
from nps_analysis.journey_config import DEFAULT_JOURNEY_CONFIG


def _sheet(workbook, spec, headers, rows):
    sheet = workbook.create_sheet(spec.sheet.name)
    for column, header in enumerate(headers, 1):
        sheet.cell(spec.sheet.header_row, column, header)
    for row_number, row in enumerate(rows, spec.sheet.data_start_row):
        for column, value in enumerate(row, 1):
            sheet.cell(row_number, column, value)


def test_build_journeys_maps_events_and_uses_touchpoint_rules(tmp_path):
    config = replace(
        DEFAULT_JOURNEY_CONFIG,
        consultation_month_range=("2026-07", "2026-07"),
        complaint_month_range=("2026-07", "2026-07"),
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    phone = "13800000000"
    _sheet(
        workbook,
        config.subscription,
        ["手机号码", "操作时间", "订购类型", "资费名称"],
        [
            [phone, "2026-07-01 10:00:00", "退订", "套餐A"],
            [phone, "2026-07-01 10:00:00", "订购", "手机邮箱"],
            [phone, "2026-07-01 10:00:00", "订购", "和彩铃"],
        ],
    )
    _sheet(
        workbook,
        config.plan_change,
        ["手机号码", "宽带号码", "办理时间", "变更标识（0：平移；1：升档；-1：降档）"],
        [[phone, "b1", 20260702101010, -1], [phone, "b1", 20260702101010, -1]],
    )
    _sheet(
        workbook,
        config.speed_package,
        ["手机号码", "办理时间", "业务类型"],
        [[phone, "2026-07-03 10:00:00", "加速包"]],
    )
    _sheet(
        workbook,
        config.wolf_number,
        ["手机号码", "狼号接触时间", "通话时长"],
        [[phone, "2026-07-04 10:00:00", 0]],
    )
    _sheet(
        workbook,
        config.complaint,
        ["手机号码", "时间", "投诉节点", "明细"],
        [
            [
                phone,
                20260705,
                "服务->移动业务->资费套餐->套餐变更",
                "客户咨询资费并要求退费",
            ]
        ],
    )
    _sheet(
        workbook,
        config.touchpoint,
        [
            "手机号码",
            "触点时间",
            "触点类型",
            "触点小类",
            "节点",
                "投诉明细",
            "投诉类型",
            "问题分类",
            "语音转文本-客户",
            "语音转文本-客服",
        ],
        [
            [
                phone,
                "2026-07-06 10:00:00",
                "资费",
                "资费咨询",
                    "宽带服务",
                    "",
                "",
                "宽带问题",
                "断网后要求退还费用",
                "",
            ],
            [
                phone,
                20260705,
                "投诉",
                "升级投诉",
                    "资费套餐",
                    "客户咨询资费并要求退费",
                "资费咨询",
                "资费咨询",
                "咨询资费",
                "",
            ],
            [
                phone,
                "2026-08-01 08:00:00",
                "资费",
                "资费咨询",
                    "资费套餐",
                    "",
                "",
                "资费问题",
                "退费",
                "",
            ],
        ],
    )
    path = tmp_path / "journey.xlsx"
    workbook.save(path)

    journey = build_journeys(path, config)[phone]

    assert [event["action"] for event in journey].count("资费变更") == 1
    assert [event["action"] for event in journey].count("办理") == 2
    assert [event["action"] for event in journey].count("投诉") == 1
    assert [event["action"] for event in journey].count("咨询") == 1
    grouped_subscription = next(event for event in journey if event["intent"] == "订购")
    plan_change = next(event for event in journey if event["action"] == "资费变更")
    speed_package = next(
        event for event in journey if event["intent"] == "流量应急服务"
    )
    consultation = next(event for event in journey if event["action"] == "咨询")
    complaint = next(event for event in journey if event["action"] == "投诉")
    assert grouped_subscription["business"] == "增值服务"
    assert grouped_subscription["intent"] == "订购"
    assert plan_change["business"] == "套餐资费"
    assert plan_change["intent"] == "降档"
    assert speed_package["business"] == "加包服务"
    assert speed_package["intent"] == "流量应急服务"
    assert consultation["business"] == "宽带服务"
    assert consultation["intent"] == "费用/资费/优惠"
    assert consultation["source"] == ""
    assert consultation["raw_text"] == ""
    assert complaint["business"] == "套餐资费"
    assert complaint["intent"] == "费用争议|退费"
    assert set(journey[0]) == {
        "event_id",
        "event_time",
        "action",
        "business",
        "intent",
        "source",
        "raw_text",
    }
